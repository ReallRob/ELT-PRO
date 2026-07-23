import importlib.metadata as metadata
import json
from pathlib import Path
from packaging.markers import default_environment
from packaging.requirements import Requirement
from packaging.utils import canonicalize_name
import shutil
import subprocess
import sys
import textwrap
import time
import unittest


PROJECT_ROOT = Path(__file__).resolve().parents[1]
EXTENSIONS_DIR = PROJECT_ROOT / "extensions"
print(EXTENSIONS_DIR)
SITE_PACKAGES_DIR = EXTENSIONS_DIR / "site-packages"
ROOT_DISTRIBUTIONS = ("pandas", "openpyxl", "selenium")


def _remove_extensions_directory():
    for _ in range(10):
        shutil.rmtree(EXTENSIONS_DIR, ignore_errors=True)
        if not EXTENSIONS_DIR.exists():
            return
        time.sleep(0.25)
    raise RuntimeError(f"Test extensions directory could not be removed: {EXTENSIONS_DIR}")


def _requirement_applies(requirement, requested_extras):
    if requirement.marker is None:
        return True

    environment = default_environment()
    if requirement.marker.evaluate(environment):
        return True

    return any(
        requirement.marker.evaluate(dict(environment, extra=extra))
        for extra in requested_extras
    )


def _distribution_closure(distribution_names):
    """Return installed distributions needed by the requested roots."""
    pending = [(name, set()) for name in distribution_names]
    requested_extras = {}
    processed_extras = {}
    distributions = {}

    while pending:
        distribution_name, extras = pending.pop()
        key = canonicalize_name(distribution_name)
        all_extras = requested_extras.setdefault(key, set())
        all_extras.update(extras)

        if key in processed_extras and all_extras <= processed_extras[key]:
            continue

        distribution = metadata.distribution(distribution_name)
        distributions[key] = distribution
        processed_extras[key] = set(all_extras)

        for requirement_text in distribution.requires or ():
            requirement = Requirement(requirement_text)
            if _requirement_applies(requirement, all_extras):
                pending.append((requirement.name, set(requirement.extras)))

    return distributions.values()


def _copy_distribution(distribution, destination):
    for relative_path in distribution.files or ():
        relative_path = Path(relative_path)
        if relative_path.is_absolute() or ".." in relative_path.parts:
            continue

        source = Path(distribution.locate_file(relative_path))
        target = destination / relative_path
        if source.is_file():
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)


class ExtensionSitePackagesTests(unittest.TestCase):
    def setUp(self):
        if EXTENSIONS_DIR.exists():
            self.skipTest("Smoke test does not modify a pre-existing extensions directory")
        SITE_PACKAGES_DIR.mkdir(parents=True)
        self.addCleanup(_remove_extensions_directory)

    def test_pandas_openpyxl_and_selenium_load_from_extensions(self):
        for distribution in _distribution_closure(ROOT_DISTRIBUTIONS):
            _copy_distribution(distribution, SITE_PACKAGES_DIR)

        script = textwrap.dedent(
            """
            import importlib.util
            import json
            from pathlib import Path
            import sys
            import tempfile

            extension_site_packages = Path(sys.argv[1]).resolve()
            roots = ("pandas", "openpyxl", "selenium")
            missing_before_extension_path = {
                name: importlib.util.find_spec(name) is None for name in roots
            }
            if not all(missing_before_extension_path.values()):
                raise AssertionError(
                    "A package was visible before the extension path was added: "
                    + repr(missing_before_extension_path)
                )

            sys.path.insert(0, str(extension_site_packages))

            import certifi
            import dateutil
            import et_xmlfile
            import numpy as np
            import openpyxl
            import pandas as pd
            import selenium
            import trio
            import urllib3
            import websocket
            from openpyxl import Workbook, load_workbook
            from selenium import webdriver
            from selenium.webdriver.common.by import By

            frame = pd.DataFrame({"value": [2, 3]})
            pandas_total = int(frame["value"].sum())

            workbook_path = Path(tempfile.mkdtemp()) / "extension-check.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "loaded from extensions"
            workbook.save(workbook_path)
            loaded = load_workbook(workbook_path, read_only=True)
            workbook_value = loaded.active["A1"].value
            loaded.close()

            modules = {
                "pandas": pd,
                "numpy": np,
                "dateutil": dateutil,
                "openpyxl": openpyxl,
                "et_xmlfile": et_xmlfile,
                "selenium": selenium,
                "certifi": certifi,
                "trio": trio,
                "urllib3": urllib3,
                "websocket": websocket,
            }

            def is_in_extensions(module):
                try:
                    Path(module.__file__).resolve().relative_to(extension_site_packages)
                    return True
                except ValueError:
                    return False

            origins = {name: str(module.__file__) for name, module in modules.items()}
            outside_extensions = [
                name for name, module in modules.items() if not is_in_extensions(module)
            ]
            if outside_extensions:
                raise AssertionError(
                    "Modules did not load from extensions: "
                    + repr({name: origins[name] for name in outside_extensions})
                )

            print(json.dumps({
                "pandas_total": pandas_total,
                "workbook_value": workbook_value,
                "selenium_by_id": By.ID,
                "webdriver_module": webdriver.__name__,
                "origins": origins,
            }))
            """
        )
        completed = subprocess.run(
            [sys.executable, "-S", "-c", script, str(SITE_PACKAGES_DIR)],
            capture_output=True,
            text=True,
        )
        if completed.returncode:
            self.fail(
                "Isolated extension import failed.\n"
                f"stdout:\n{completed.stdout}\n"
                f"stderr:\n{completed.stderr}"
            )
        result = json.loads(completed.stdout)

        self.assertEqual(result["pandas_total"], 5)
        self.assertEqual(result["workbook_value"], "loaded from extensions")
        self.assertEqual(result["selenium_by_id"], "id")
        self.assertEqual(result["webdriver_module"], "selenium.webdriver")

        launcher = subprocess.run(
            [
                sys.executable,
                "-S",
                str(PROJECT_ROOT / "tests" / "extensions_main.py"),
                str(SITE_PACKAGES_DIR),
            ],
            capture_output=True,
            text=True,
        )
        if launcher.returncode:
            self.fail(
                "Extensions main entry point failed.\n"
                f"stdout:\n{launcher.stdout}\n"
                f"stderr:\n{launcher.stderr}"
            )
        self.assertIn("value", launcher.stdout)
        self.assertIn("2", launcher.stdout)
        self.assertIn("3", launcher.stdout)


if __name__ == "__main__":
    unittest.main()
