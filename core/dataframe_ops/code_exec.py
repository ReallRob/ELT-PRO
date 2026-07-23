"""Shared execution helpers for user-written pandas/openpyxl code."""

from __future__ import annotations

import ast
import builtins as py_builtins
from contextlib import redirect_stderr, redirect_stdout
from copy import copy, deepcopy
from dataclasses import dataclass
from datetime import date, datetime, timedelta
import keyword
import math
import multiprocessing
from queue import Empty
import re
import time
from types import ModuleType
import traceback

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from parameter_resolver import (
    normalize_parameter_mappings,
    normalize_runtime_parameters,
    resolve_placeholder,
)
from core.runtime_extensions import activate_external_extensions

DEFAULT_CODE_TIMEOUT_SECONDS = 10
MAX_CODE_TIMEOUT_SECONDS = 3600
MAX_LOG_LINE_CHARS = 10000
PROCESS_CANCEL_GRACE_SECONDS = 0.75
_LOCAL_RETURN_KEY = "__code_block_locals__"
PRESET_IMPORT_SNIPPET = """import pandas as pd
import numpy as np
import re
import math
import openpyxl
from datetime import datetime, date, timedelta
from copy import copy, deepcopy
from openpyxl.utils import get_column_letter"""


@dataclass
class CodeExecutionOutput:
    name: str
    data: object
    data_type: str = "table"


@dataclass
class CodeExecutionResult:
    outputs: list[CodeExecutionOutput]
    state: dict


_ALIAS_RE = re.compile(r"^[A-Za-z_]\w*$")
_RESERVED_ALIASES = {
    "pd",
    "np",
    "re",
    "math",
    "datetime",
    "date",
    "timedelta",
    "openpyxl",
    "copy",
    "deepcopy",
    "get_column_letter",
    "params",
    "mappings",
    "param",
    "state",
    "dfs",
    "wbs",
    "wss",
    "locals",
    "should_cancel",
    "check_cancel",
    "__builtins__",
    "__build_class__",
    "__name__",
}
_ALLOWED_WRAPPED_NAMES = {
    "pd",
    "np",
    "re",
    "math",
    "datetime",
    "date",
    "timedelta",
    "openpyxl",
    "copy",
    "deepcopy",
    "get_column_letter",
    "params",
    "mappings",
    "param",
    "state",
    "dfs",
    "df",
    "wbs",
    "wss",
    "wb",
    "ws",
    "should_cancel",
    "check_cancel",
}


def is_valid_code_alias(alias):
    text = str(alias or "").strip()
    return bool(
        text
        and _ALIAS_RE.match(text)
        and not keyword.iskeyword(text)
        and text not in _RESERVED_ALIASES
    )


def _is_valid_function_space_module_name(name):
    text = str(name or "").strip()
    return is_valid_code_alias(text)


def coerce_timeout_seconds(value):
    try:
        seconds = float(value)
    except (TypeError, ValueError):
        seconds = DEFAULT_CODE_TIMEOUT_SECONDS
    return max(1.0, min(float(MAX_CODE_TIMEOUT_SECONDS), seconds))


class _CallbackWriter:
    def __init__(self, log_callback, stream):
        self._log_callback = log_callback
        self._stream = stream
        self._buffer = ""

    def write(self, text):
        value = str(text or "")
        if not value:
            return 0
        self._buffer += value
        while "\n" in self._buffer:
            line, self._buffer = self._buffer.split("\n", 1)
            self._emit(line.rstrip("\r"))
        if len(self._buffer) > MAX_LOG_LINE_CHARS:
            self._emit(self._buffer[:MAX_LOG_LINE_CHARS] + " ...[已截断]")
            self._buffer = ""
        return len(value)

    def flush(self):
        if self._buffer:
            self._emit(self._buffer.rstrip("\r"))
            self._buffer = ""

    def _emit(self, line):
        _emit_code_log(self._log_callback, self._stream, str(line))


def _emit_code_log(log_callback, stream, message):
    if not callable(log_callback):
        return
    if stream == "stderr":
        prefix = "代码块错误输出"
    elif stream == "info":
        prefix = "代码块提示"
    else:
        prefix = "代码块输出"
    text = f"{prefix}: {message}" if str(message) else f"{prefix}:"
    try:
        log_callback(text)
    except Exception:
        pass


def _code_block_import(name, globals=None, locals=None, fromlist=(), level=0, module_registry=None):
    """Resolve virtual function modules, then use Python's normal import behavior."""
    if level != 0:
        raise ImportError("代码块不支持相对 import")
    module_name = str(name or "").strip()
    root = module_name.split(".", 1)[0]
    virtual_modules = module_registry or {}
    if root in virtual_modules:
        if module_name != root:
            raise ImportError(f"函数空间引用 {root!r} 不支持子模块 import")
        return virtual_modules[root]
    return py_builtins.__import__(name, globals, locals, fromlist, level)


def _build_code_builtins(module_registry=None):
    def code_block_import(name, globals=None, locals=None, fromlist=(), level=0):
        return _code_block_import(name, globals, locals, fromlist, level, module_registry)

    return {
        "__import__": code_block_import,
        "__build_class__": py_builtins.__build_class__,
        "abs": abs,
        "all": all,
        "any": any,
        "bool": bool,
        "callable": callable,
        "classmethod": classmethod,
        "dict": dict,
        "enumerate": enumerate,
        "Exception": Exception,
        "float": float,
        "getattr": getattr,
        "hasattr": hasattr,
        "int": int,
        "isinstance": isinstance,
        "len": len,
        "list": list,
        "locals": locals,
        "max": max,
        "min": min,
        "object": object,
        "open": py_builtins.open,
        "print": print,
        "property": property,
        "range": range,
        "repr": repr,
        "round": round,
        "set": set,
        "setattr": setattr,
        "staticmethod": staticmethod,
        "sorted": sorted,
        "str": str,
        "sum": sum,
        "super": super,
        "tuple": tuple,
        "type": type,
        "zip": zip,
    }


def _prepare_tables(raw_tables):
    tables = {}
    for alias, df in (raw_tables or {}).items():
        alias = str(alias or "").strip()
        if not alias:
            continue
        if not is_valid_code_alias(alias):
            raise ValueError(f"变量名无效: {alias}")
        if not isinstance(df, pd.DataFrame):
            raise TypeError(f"变量 {alias} 对应的输入不是 DataFrame")
        tables[alias] = df.copy()
    return tables


def _prepare_workbooks(raw_workbooks):
    workbooks = {}
    workbook_entries = {}
    for alias, value in (raw_workbooks or {}).items():
        alias = str(alias or "").strip()
        if not alias:
            continue
        if not is_valid_code_alias(alias):
            raise ValueError(f"变量名无效: {alias}")
        wb = value.get("_wb") if isinstance(value, dict) else value
        if not isinstance(wb, Workbook):
            raise TypeError(f"变量 {alias} 对应的输入不是 Workbook")
        workbooks[alias] = wb
        workbook_entries[alias] = value if isinstance(value, dict) else {"_wb": wb}
    return workbooks, workbook_entries


def _prepare_worksheets(raw_worksheets, workbooks):
    worksheets = {}
    for alias, item in (raw_worksheets or {}).items():
        alias = str(alias or "").strip()
        if not alias:
            continue
        if not is_valid_code_alias(alias):
            raise ValueError(f"变量名无效: {alias}")
        wb_alias = str((item or {}).get("workbook_alias") or "").strip()
        sheet_name = str((item or {}).get("sheet_name") or "").strip()
        wb = workbooks.get(wb_alias)
        if wb is None:
            continue
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"工作簿 {wb_alias} 不存在工作表: {sheet_name}")
            worksheets[alias] = wb[sheet_name]
        else:
            worksheets[alias] = wb.active
    return worksheets


def _ensure_unique_aliases(*groups):
    seen = set()
    for group in groups:
        for alias in group:
            if alias in seen:
                raise ValueError(f"变量名重复: {alias}")
            seen.add(alias)


def _normalize_function_spaces(raw_spaces, legacy_global_code=""):
    spaces = []
    for index, item in enumerate(raw_spaces or [], start=1):
        if not isinstance(item, dict):
            continue
        code = str(item.get("code") or "")
        namespace = str(item.get("namespace") or item.get("id") or f"funcs{index}").strip()
        if not namespace:
            namespace = f"funcs{index}"
        if not _is_valid_function_space_module_name(namespace):
            raise ValueError(f"函数空间引用名称无效: {namespace}")
        spaces.append(
            {
                "id": str(item.get("id") or namespace),
                "name": str(item.get("name") or namespace),
                "namespace": namespace,
                "enabled": bool(item.get("enabled", True)),
                "expose_globals": bool(item.get("expose_globals", False)),
                "code": code,
            }
        )
    if not spaces and str(legacy_global_code or "").strip():
        spaces.append(
            {
                "id": "legacy_global",
                "name": "旧全局函数",
                "namespace": "global_funcs",
                "enabled": True,
                "expose_globals": True,
                "code": str(legacy_global_code or ""),
            }
        )
    return spaces


def _public_space_values(space_env, hidden_names=None):
    hidden = set(_ALLOWED_WRAPPED_NAMES) | _RESERVED_ALIASES | {"__builtins__"}
    hidden.update(hidden_names or ())
    public = {}
    for name, value in (space_env or {}).items():
        if not is_valid_code_alias(name) or name in hidden:
            continue
        public[name] = value
    return public


def _validate_function_space_reference_names(spaces):
    owners = {}
    for space in spaces:
        if not space.get("enabled", True):
            continue
        label = str(space.get("name") or space.get("namespace") or "函数空间")
        reference_name = str(space.get("namespace") or "").strip()
        existing = owners.get(reference_name)
        if existing is not None:
            raise ValueError(f"函数空间引用名称重复: {reference_name}（{existing}、{label}）")
        owners[reference_name] = label


def _create_function_space_module(namespace):
    module = ModuleType(namespace)
    module.__all__ = ()
    return module


def _function_space_dependencies(space, reference_names):
    code = str(space.get("code") or "")
    if not code.strip():
        return set()
    try:
        tree = ast.parse(code, mode="exec")
    except SyntaxError:
        return set()

    namespace = str(space.get("namespace") or "").strip()
    dependencies = set()
    for node in tree.body:
        if not isinstance(node, ast.ImportFrom) or node.level:
            continue
        reference_name = str(node.module or "").split(".", 1)[0]
        if reference_name == namespace:
            raise ValueError(f"函数空间不能导入自身: {namespace}")
        if reference_name in reference_names:
            dependencies.add(reference_name)
    return dependencies


def _function_space_execution_order(spaces):
    active_spaces = [space for space in spaces if space.get("enabled", True)]
    reference_names = {str(space.get("namespace") or "").strip() for space in active_spaces}
    dependencies = {
        str(space.get("namespace") or "").strip(): _function_space_dependencies(space, reference_names)
        for space in active_spaces
    }
    pending = list(active_spaces)
    executed = set()
    ordered = []

    while pending:
        ready = [
            space
            for space in pending
            if dependencies[str(space.get("namespace") or "").strip()].issubset(executed)
        ]
        if not ready:
            names = [str(space.get("namespace") or "").strip() for space in pending]
            raise ValueError(
                "函数空间存在循环引用: "
                + " -> ".join(names)
                + "；请改用 import 引用名称，并在函数体内调用对方函数"
            )
        for space in ready:
            namespace = str(space.get("namespace") or "").strip()
            ordered.append(space)
            executed.add(namespace)
            pending.remove(space)
    return ordered


def _execute_function_spaces(exec_env, function_spaces, legacy_global_code="", module_registry=None):
    spaces = _normalize_function_spaces(function_spaces, legacy_global_code)
    _validate_function_space_reference_names(spaces)
    registry = module_registry if module_registry is not None else {}
    ordered_spaces = _function_space_execution_order(spaces)
    modules = {}
    for space in ordered_spaces:
        namespace = str(space.get("namespace") or "").strip()
        module = _create_function_space_module(namespace)
        modules[namespace] = module
        registry[namespace] = module
        exec_env[namespace] = module

    for space in ordered_spaces:
        code = str(space.get("code") or "")
        namespace = str(space.get("namespace") or "").strip()
        module = modules[namespace]
        for name, value in exec_env.items():
            module.__dict__.setdefault(name, value)
        module.__dict__["__name__"] = namespace
        if code.strip():
            exec(compile(code, f"<函数空间:{namespace}>", "exec"), module.__dict__, module.__dict__)
        public = _public_space_values(module.__dict__, modules)
        module.__all__ = tuple(name for name in public if not name.startswith("_"))
        if space.get("expose_globals", False):
            exec_env.update(public)
    return spaces


def _fallback_output_name(output_names, index):
    if index - 1 < len(output_names):
        name = str(output_names[index - 1] or "").strip()
        if name:
            return name
    return "代码块结果" if index == 1 else f"代码块结果{index}"


def _is_workbook_entry(value):
    return isinstance(value, dict) and isinstance(value.get("_wb"), Workbook)


def _workbook_entry_for_workbook(wb, workbook_entries):
    for entry in (workbook_entries or {}).values():
        if _is_workbook_entry(entry) and entry.get("_wb") is wb:
            return entry
    return {"_wb": wb}


def _is_formal_output_value(value):
    if isinstance(value, pd.DataFrame):
        return True
    if _is_workbook_entry(value):
        return True
    return isinstance(value, (Workbook, Worksheet))


def _ordinary_value_to_dataframe(value):
    if isinstance(value, pd.Series):
        return value.to_frame(name=value.name or "值")
    if isinstance(value, dict):
        return pd.DataFrame([value])
    if isinstance(value, (list, tuple)):
        items = list(value)
        if not items:
            return pd.DataFrame(columns=["值"])
        if all(isinstance(item, dict) for item in items):
            return pd.DataFrame(items)
        if all(isinstance(item, (list, tuple)) for item in items):
            return pd.DataFrame(items)
        return pd.DataFrame({"值": items})
    return pd.DataFrame([{"值": value}])


def _make_code_output(name, value, output_names, index, workbook_entries=None):
    output_name = str(name or "").strip() or _fallback_output_name(output_names, index)
    if isinstance(value, pd.DataFrame):
        return {"name": output_name, "data_type": "table", "data": value}
    if _is_workbook_entry(value):
        return {"name": output_name, "data_type": "workbook", "data": value}
    if isinstance(value, Workbook):
        entry = _workbook_entry_for_workbook(value, workbook_entries)
        return {
            "name": output_name,
            "data_type": "workbook",
            "data": entry,
        }
    if isinstance(value, Worksheet):
        wb = getattr(value, "parent", None)
        if isinstance(wb, Workbook):
            entry = _workbook_entry_for_workbook(wb, workbook_entries)
            return {
                "name": output_name,
                "data_type": "workbook",
                "data": entry,
            }
    return {
        "name": output_name,
        "data_type": "table",
        "data": _ordinary_value_to_dataframe(value),
    }


def _normalize_code_outputs(raw_result, output_names, workbook_entries=None):
    if _is_workbook_entry(raw_result):
        return [_make_code_output("", raw_result, output_names, 1, workbook_entries)]

    if isinstance(raw_result, dict):
        output_flags = [_is_formal_output_value(value) for value in raw_result.values()]
        if output_flags and all(output_flags):
            outputs = []
            for index, (name, value) in enumerate(raw_result.items(), start=1):
                output_name = str(name or "").strip()
                if not output_name:
                    raise ValueError("代码块返回字典的 key 必须是非空输出名称")
                outputs.append(
                    _make_code_output(output_name, value, output_names, index, workbook_entries)
                )
            return outputs
        if any(output_flags):
            raise TypeError(
                "代码块返回字典不能混合 DataFrame/Workbook/Worksheet 和普通值；"
                "请把普通值包装成 DataFrame 后再作为一个输出。"
            )
        return [_make_code_output("", raw_result, output_names, 1, workbook_entries)]

    if isinstance(raw_result, (list, tuple)):
        output_flags = [_is_formal_output_value(value) for value in raw_result]
        if output_flags and all(output_flags):
            return [
                _make_code_output("", value, output_names, index, workbook_entries)
                for index, value in enumerate(raw_result, start=1)
            ]
        if any(output_flags):
            raise TypeError(
                "代码块返回列表/元组不能混合 DataFrame/Workbook/Worksheet 和普通值；"
                "请把普通值包装成 DataFrame 后再作为一个输出。"
            )
        return [_make_code_output("", raw_result, output_names, 1, workbook_entries)]

    return [_make_code_output("", raw_result, output_names, 1, workbook_entries)]


def _wrap_code_as_main(code, available_names):
    user_args = ", ".join(
        f"{name}={name}"
        for name in available_names
        if is_valid_code_alias(name) or name in _ALLOWED_WRAPPED_NAMES
    )
    hidden_arg = "__code_block_locals_fn=__builtins__['locals']"
    args = f"{user_args}, {hidden_arg}" if user_args else hidden_arg
    body_lines = str(code or "").splitlines()
    indented = [f"    {line}" if line.strip() else "" for line in body_lines]
    if not indented:
        indented = ["    pass"]
    indented.append(f"    return {{{_LOCAL_RETURN_KEY!r}: __code_block_locals_fn()}}")
    return f"def __code_block_main__({args}):\n" + "\n".join(indented)


def _extract_leading_star_imports(code, exec_env):
    """Run leading ``from module import *`` statements before function wrapping."""
    source = str(code or "")
    if "import *" not in source:
        return source
    try:
        tree = ast.parse(source, mode="exec")
    except SyntaxError:
        return source

    lines = source.splitlines()
    imports = []
    seen_non_import = False
    for node in tree.body:
        is_import = isinstance(node, (ast.Import, ast.ImportFrom))
        is_star_import = isinstance(node, ast.ImportFrom) and any(
            item.name == "*" for item in node.names
        )
        if is_star_import:
            if seen_non_import:
                raise SyntaxError("from 模块 import * 只能放在当前代码的顶部")
            if node.lineno != node.end_lineno:
                raise SyntaxError("from 模块 import * 必须单独写在一行")
            trailing = lines[node.end_lineno - 1][node.end_col_offset:].strip()
            if trailing and not trailing.startswith("#"):
                raise SyntaxError("from 模块 import * 必须单独写在一行")
            imports.append(node)
        if not is_import:
            seen_non_import = True

    for node in imports:
        statement = ast.get_source_segment(source, node)
        exec(compile(statement, "<当前代码块导入>", "exec"), exec_env, exec_env)
        for line_number in range(node.lineno - 1, node.end_lineno):
            lines[line_number] = ""
    return "\n".join(lines)


def _is_state_mapping(value):
    return isinstance(value, dict) or (
        hasattr(value, "get") and hasattr(value, "items") and hasattr(value, "__setitem__")
    )


def _execute_code_payload(payload):
    activate_external_extensions()
    tables = _prepare_tables(payload.get("tables") or {})
    raw_workbooks = payload.get("workbooks") or {}
    workbooks, workbook_entries = _prepare_workbooks(raw_workbooks)
    worksheets = _prepare_worksheets(payload.get("worksheets") or {}, workbooks)
    _ensure_unique_aliases(tables, workbooks, worksheets)
    params = normalize_runtime_parameters(payload.get("runtime_parameters") or {})
    mappings = normalize_parameter_mappings(payload.get("parameter_mappings") or {})
    raw_state = payload.get("state")
    state = raw_state if _is_state_mapping(raw_state) else {}
    fallback_alias = str(payload.get("fallback_alias") or "df").strip() or "df"

    def should_cancel():
        return str(state.get("run_status") or "").strip().lower() == "stopped"

    def check_cancel():
        if should_cancel():
            raise RuntimeError("代码块已停止")
        return False

    def param(name, mapping_name=None):
        expression = str(name).strip()
        if mapping_name:
            expression = f"{expression}|map:{str(mapping_name).strip()}"
        return resolve_placeholder(expression, params, mappings, strict=True)

    dfs = dict(tables)
    wbs = dict(workbooks)
    wss = dict(worksheets)
    primary_df = tables.get(fallback_alias)
    if primary_df is None:
        primary_df = next(iter(tables.values())) if tables else None
    primary_wb_alias = "wb" if "wb" in workbooks else (next(iter(workbooks.keys())) if workbooks else "")
    primary_wb = workbooks.get(primary_wb_alias)
    primary_ws_alias = "ws" if "ws" in worksheets else (next(iter(worksheets.keys())) if worksheets else "")
    primary_ws = worksheets.get(primary_ws_alias)

    module_registry = {}
    exec_env = {
        "__builtins__": _build_code_builtins(module_registry),
        "__name__": str(payload.get("module_name") or "__code_block__"),
        "pd": pd,
        "np": np,
        "re": re,
        "math": math,
        "datetime": datetime,
        "date": date,
        "timedelta": timedelta,
        "openpyxl": openpyxl,
        "copy": copy,
        "deepcopy": deepcopy,
        "get_column_letter": get_column_letter,
        "params": params,
        "mappings": mappings,
        "param": param,
        "state": state,
        "dfs": dfs,
        "df": primary_df,
        "wbs": wbs,
        "wss": wss,
        "wb": primary_wb,
        "should_cancel": should_cancel,
        "check_cancel": check_cancel,
    }
    if primary_ws is not None:
        exec_env["ws"] = primary_ws
    exec_env.update(tables)
    exec_env.update(workbooks)
    exec_env.update(worksheets)

    global_code = str(payload.get("global_code") or "")
    function_spaces = payload.get("function_spaces") or []
    node_code = str(payload.get("code") or "")
    _execute_function_spaces(exec_env, function_spaces, global_code, module_registry)

    explicit_return = False
    returned = None
    local_vars = {}
    if node_code.strip():
        prepared_code = _extract_leading_star_imports(node_code, exec_env)
        wrapped = _wrap_code_as_main(prepared_code, exec_env.keys())
        exec(compile(wrapped, "<当前代码块>", "exec"), exec_env, exec_env)
        try:
            returned = exec_env["__code_block_main__"]()
        except SystemExit as exc:
            if not payload.get("allow_system_exit", False):
                raise
            if exc.code not in (None, 0):
                raise RuntimeError(f"独立进程以非零状态退出: {exc.code}")
            returned = None
        if isinstance(returned, dict) and set(returned.keys()) == {_LOCAL_RETURN_KEY}:
            local_vars = dict(returned.get(_LOCAL_RETURN_KEY) or {})
        else:
            explicit_return = True

    target_col = str(payload.get("target_col") or "").strip()
    if fallback_alias in local_vars:
        fallback_df = local_vars[fallback_alias]
    elif fallback_alias in exec_env:
        fallback_df = exec_env[fallback_alias]
    elif "df" in local_vars:
        fallback_df = local_vars["df"]
    elif "df" in exec_env:
        fallback_df = exec_env["df"]
    else:
        fallback_df = primary_df

    target_source = local_vars if target_col in local_vars else exec_env
    if target_col and target_col in target_source:
        if not isinstance(fallback_df, pd.DataFrame):
            raise TypeError(f"变量 {fallback_alias} 必须是 DataFrame，才能写入输出列")
        fallback_df[target_col] = target_source[target_col]

    if explicit_return:
        raw_result = returned
    else:
        raw_result = None

    outputs = [] if raw_result is None else _normalize_code_outputs(
        raw_result,
        payload.get("output_names") or [],
        workbook_entries,
    )
    return {"outputs": outputs, "state": dict(state)}


def _outputs_from_items(items):
    return [
        CodeExecutionOutput(
            name=str(item.get("name") or f"代码结果{index}"),
            data=item.get("data"),
            data_type=str(item.get("data_type") or "table"),
        )
        for index, item in enumerate(items or [], start=1)
        if isinstance(item, dict)
    ]


def _run_dataframe_code_inline(payload, log_callback=None):
    payload = dict(payload)
    if callable(log_callback):
        stdout_writer = _CallbackWriter(log_callback, "stdout")
        stderr_writer = _CallbackWriter(log_callback, "stderr")
        try:
            with redirect_stdout(stdout_writer), redirect_stderr(stderr_writer):
                result = _execute_code_payload(payload)
        except Exception:
            raise ValueError(f"{payload.get('error_prefix') or '代码执行'}失败：\n{traceback.format_exc()}")
        finally:
            stdout_writer.flush()
            stderr_writer.flush()
    else:
        try:
            result = _execute_code_payload(payload)
        except Exception:
            raise ValueError(f"{payload.get('error_prefix') or '代码执行'}失败：\n{traceback.format_exc()}")
    return CodeExecutionResult(
        outputs=_outputs_from_items(result.get("outputs") or []),
        state=result.get("state") if isinstance(result.get("state"), dict) else {},
    )


def _code_process_worker(payload, event_queue):
    def log_callback(message):
        event_queue.put(("log", str(message)))

    try:
        result = _run_dataframe_code_inline(payload, log_callback)
    except BaseException:
        event_queue.put(("error", traceback.format_exc()))
    else:
        event_queue.put(("result", result))


def _process_cancel_requested(payload):
    state = payload.get("state")
    return isinstance(state, dict) and str(state.get("run_status") or "").strip().lower() == "stopped"


def _sync_process_state(source, target):
    if not isinstance(source, dict):
        return
    for key, value in source.items():
        target[key] = value


def _run_dataframe_code_in_process(payload, log_callback=None):
    if payload.get("workbooks") or payload.get("worksheets"):
        raise ValueError("独立进程运行不支持内存 Workbook 输入，请使用普通运行模式")

    child_payload = dict(payload)
    child_payload["module_name"] = "__main__"
    child_payload["allow_system_exit"] = True
    context = multiprocessing.get_context("spawn")
    manager = context.Manager()
    shared_state = manager.dict()
    _sync_process_state(payload.get("state"), shared_state)
    child_payload["state"] = shared_state
    event_queue = context.Queue()
    process = context.Process(target=_code_process_worker, args=(child_payload, event_queue))
    result = None
    failure = ""
    cancel_started_at = None

    try:
        try:
            process.start()
        except Exception as exc:
            raise ValueError(f"独立进程启动失败: {exc}") from exc

        while result is None and not failure:
            _sync_process_state(payload.get("state"), shared_state)
            if _process_cancel_requested(payload):
                shared_state["run_status"] = "stopped"
                if cancel_started_at is None:
                    cancel_started_at = time.monotonic()
                    _emit_code_log(log_callback, "info", "已同步停止状态，等待独立进程协作退出。")
                elif time.monotonic() - cancel_started_at >= PROCESS_CANCEL_GRACE_SECONDS:
                    if process.is_alive():
                        process.terminate()
                    process.join()
                    raise RuntimeError("代码块已停止，独立进程已终止")
            try:
                event_type, event_value = event_queue.get(timeout=0.1)
            except Empty:
                if process.is_alive():
                    continue
                try:
                    event_type, event_value = event_queue.get(timeout=0.2)
                except Empty:
                    break
            if event_type == "log":
                if callable(log_callback):
                    try:
                        log_callback(str(event_value))
                    except Exception:
                        pass
            elif event_type == "result":
                result = event_value
            elif event_type == "error":
                failure = str(event_value or "")

        process.join()
        if _process_cancel_requested(payload):
            raise RuntimeError("代码块已停止，独立进程已终止")
        if result is not None:
            return result
        if failure:
            raise ValueError(f"{payload.get('error_prefix') or '代码执行'}失败：\n{failure}")
        exit_code = process.exitcode
        raise ValueError(f"独立进程未返回结果，退出码: {exit_code}")
    finally:
        if process.is_alive():
            process.terminate()
            process.join()
        event_queue.close()
        event_queue.join_thread()
        manager.shutdown()


def run_dataframe_code(
    tables,
    code,
    workbooks=None,
    worksheets=None,
    global_code="",
    function_spaces=None,
    state=None,
    runtime_parameters=None,
    parameter_mappings=None,
    timeout_seconds=DEFAULT_CODE_TIMEOUT_SECONDS,
    result_name="result",
    fallback_alias="df",
    target_col="",
    output_names=None,
    error_prefix="代码执行",
    log_callback=None,
    execution_mode="auto",
):
    """Run user code inline or in an isolated child process."""
    timeout_seconds = coerce_timeout_seconds(timeout_seconds)
    runtime_state = state if isinstance(state, dict) else {}
    runtime_state.setdefault("run_status", "running")
    payload = {
        "tables": tables or {},
        "workbooks": workbooks or {},
        "worksheets": worksheets or {},
        "code": code,
        "global_code": global_code or "",
        "function_spaces": function_spaces or [],
        "state": runtime_state,
        "runtime_parameters": runtime_parameters or {},
        "parameter_mappings": parameter_mappings or {},
        "result_name": result_name,
        "fallback_alias": fallback_alias,
        "target_col": target_col,
        "output_names": output_names or [],
        "error_prefix": error_prefix,
        "timeout_seconds": timeout_seconds,
    }
    if str(execution_mode or "").strip().lower() == "process":
        _emit_code_log(log_callback, "info", "代码块正在独立进程运行；关闭子窗口后进程会结束。")
        return _run_dataframe_code_in_process(payload, log_callback)
    _emit_code_log(
        log_callback,
        "info",
        "代码块在线程内执行；停止按钮会将 state['run_status'] 设为 'stopped'，请在代码中检查该状态并自行退出。"
    )
    return _run_dataframe_code_inline(payload, log_callback)
