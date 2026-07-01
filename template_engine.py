import os
from copy import copy
from datetime import date, datetime
from numbers import Real

import pandas as pd

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill
    from openpyxl.utils import column_index_from_string
except ImportError:
    openpyxl = None
    column_index_from_string = None


TEMPLATE_PREVIEW_MAX_ROWS = 5000
TEMPLATE_PREVIEW_MAX_COLS = 200


def _coerce_preview_bound(value, default, minimum=1, allow_column_letters=False):
    if value is None or value == "":
        return default
    text = str(value).strip().replace("$", "")
    if allow_column_letters and text.isalpha() and column_index_from_string is not None:
        return max(minimum, column_index_from_string(text.upper()))
    try:
        result = int(text)
    except (TypeError, ValueError):
        return default
    return max(minimum, result)


def worksheet_to_preview_df(
    ws,
    start_row=1,
    start_col=1,
    max_rows=TEMPLATE_PREVIEW_MAX_ROWS,
    max_cols=TEMPLATE_PREVIEW_MAX_COLS,
):
    """Build a bounded preview DataFrame for one worksheet range."""
    source_rows = ws.max_row or 0
    source_cols = ws.max_column or 0
    start_row = _coerce_preview_bound(start_row, 1)
    start_col = _coerce_preview_bound(start_col, 1, allow_column_letters=True)
    max_rows = _coerce_preview_bound(max_rows, TEMPLATE_PREVIEW_MAX_ROWS, minimum=0)
    max_cols = _coerce_preview_bound(max_cols, TEMPLATE_PREVIEW_MAX_COLS, minimum=0)
    if max_rows <= 0 or max_cols <= 0 or start_row > source_rows or start_col > source_cols:
        preview_df = pd.DataFrame()
        preview_df.attrs["_hide_column_names"] = True
        preview_df.attrs["_source_rows"] = source_rows
        preview_df.attrs["_source_cols"] = source_cols
        preview_df.attrs["_preview_limited"] = bool(source_rows or source_cols)
        preview_df.attrs["_range_start_row"] = start_row
        preview_df.attrs["_range_start_col"] = start_col
        return preview_df

    end_row = min(source_rows, start_row + max_rows - 1)
    end_col = min(source_cols, start_col + max_cols - 1)
    rows = list(
        ws.iter_rows(
            min_row=start_row,
            max_row=end_row,
            min_col=start_col,
            max_col=end_col,
            values_only=True,
        )
    )
    if not rows:
        preview_df = pd.DataFrame()
    else:
        max_cols_seen = max((len(row) for row in rows), default=0)
        normalized_rows = [
            list(row) + [None] * (max_cols_seen - len(row))
            for row in rows
        ]
        columns = [f"{start_col + i}" for i in range(max_cols_seen)]
        preview_df = pd.DataFrame(normalized_rows, columns=columns)
        preview_df.index = range(start_row, start_row + len(preview_df))
    preview_df.attrs["_hide_column_names"] = True
    preview_df.attrs["_source_rows"] = source_rows
    preview_df.attrs["_source_cols"] = source_cols
    preview_df.attrs["_preview_limited"] = (
        end_row < source_rows or end_col < source_cols or start_row > 1 or start_col > 1
    )
    preview_df.attrs["_range_start_row"] = start_row
    preview_df.attrs["_range_start_col"] = start_col
    preview_df.attrs["_range_end_row"] = end_row
    preview_df.attrs["_range_end_col"] = end_col
    return preview_df


def workbook_sheet_preview_data(
    wb,
    sheet_name=None,
    saved_path="",
    start_row=1,
    start_col=1,
    max_rows=TEMPLATE_PREVIEW_MAX_ROWS,
    max_cols=TEMPLATE_PREVIEW_MAX_COLS,
):
    """Build preview payload for one worksheet and one bounded range."""
    if sheet_name is None or sheet_name == "":
        if not wb.worksheets:
            raise ValueError("workbook has no worksheets")
        ws = wb.worksheets[0]
    else:
        ws = wb[sheet_name]
    preview_df = worksheet_to_preview_df(ws, start_row, start_col, max_rows, max_cols)
    return {
        "_template_preview": True,
        "saved_path": saved_path,
        "sheets": {ws.title: preview_df},
        "sheet_count": len(wb.worksheets),
        "previewed_sheet_count": 1,
        "skipped_sheets": [sheet.title for sheet in wb.worksheets if sheet.title != ws.title],
        "preview_truncated": bool(preview_df.attrs.get("_preview_limited")),
        "preview_limit": {
            "max_rows": max_rows,
            "max_cols": max_cols,
            "max_sheets": 1,
            "start_row": start_row,
            "start_col": start_col,
        },
    }


def build_template_metadata(wb, file_path=""):
    """根据 openpyxl Workbook 生成模板结构信息。"""
    sheets_info = {}
    for name in wb.sheetnames:
        ws = wb[name]
        sheets_info[name] = {
            "name": name,
            "max_row": ws.max_row or 1,
            "max_col": ws.max_column or 1,
        }

    return {
        "path": file_path,
        "sheets": sheets_info,
        "sheet_names": list(wb.sheetnames),
    }


def workbook_to_preview_data(
    wb,
    saved_path="",
    max_sheets=None,
    start_row=1,
    start_col=1,
    max_rows=TEMPLATE_PREVIEW_MAX_ROWS,
    max_cols=TEMPLATE_PREVIEW_MAX_COLS,
):
    """Convert an openpyxl workbook to a bounded multi-sheet preview payload."""
    sheets = {}
    worksheets = list(wb.worksheets)
    if max_sheets is None:
        preview_worksheets = worksheets
    else:
        preview_worksheets = worksheets[: max(0, int(max_sheets))]
    skipped_sheets = [ws.title for ws in worksheets[len(preview_worksheets):]]
    for ws in preview_worksheets:
        sheets[ws.title] = worksheet_to_preview_df(ws, start_row, start_col, max_rows, max_cols)
    return {
        "_template_preview": True,
        "saved_path": saved_path,
        "sheets": sheets,
        "sheet_count": len(worksheets),
        "previewed_sheet_count": len(preview_worksheets),
        "skipped_sheets": skipped_sheets,
        "preview_truncated": bool(skipped_sheets),
        "preview_limit": {
            "max_rows": max_rows,
            "max_cols": max_cols,
            "max_sheets": max_sheets,
            "start_row": start_row,
            "start_col": start_col,
        },
    }




def close_workbook(wb):
    if wb is None:
        return
    close = getattr(wb, "close", None)
    if callable(close):
        try:
            close()
        except Exception:
            pass


def clone_workbook(wb):
    """Return an independent openpyxl workbook copy for safe preview/runtime edits."""
    if wb is None:
        return None
    if openpyxl is None:
        raise ImportError("缺少 openpyxl 依赖，请执行: pip install openpyxl")
    from io import BytesIO

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return openpyxl.load_workbook(stream)


def load_template(file_path):
    """
    加载模板文件，返回 (workbook, metadata_dict)。

    metadata 结构:
    {
        "path": "D:/templates/月报.xlsx",
        "sheets": {
            "Sheet1": {"name": "Sheet1", "max_row": 50, "max_col": 10},
            "Sheet2": {"name": "Sheet2", "max_row": 30, "max_col": 8}
        },
        "sheet_names": ["Sheet1", "Sheet2"]
    }
    """
    if openpyxl is None:
        raise ImportError("缺少 openpyxl 依赖，请执行: pip install openpyxl")

    wb = openpyxl.load_workbook(file_path)
    return wb, build_template_metadata(wb, file_path)


def _is_blank(expr):
    return expr is None or str(expr).strip() == ""


def _eval_position(expr, max_row, max_col, label="位置", allow_column_letters=False):
    """安全计算位置表达式，可用变量: max_row, max_col。"""
    if isinstance(expr, (int, float)):
        return int(expr)

    text = str(expr).strip()
    if not text:
        raise ValueError(f"{label}不能为空")

    if allow_column_letters:
        column_text = text.replace("$", "")
        if column_text.isalpha() and column_index_from_string is not None:
            return column_index_from_string(column_text.upper())

    allowed = {"max_row": max_row, "max_col": max_col}
    try:
        result = eval(text, {"__builtins__": {}}, allowed)
        return int(result)
    except Exception:
        raise ValueError(
            f"无法解析{label}表达式: {expr} "
            f"(可用变量: max_row={max_row}, max_col={max_col})"
        )


def _eval_optional_position(expr, max_row, max_col, label="位置", allow_column_letters=False):
    if _is_blank(expr):
        return None
    return _eval_position(expr, max_row, max_col, label, allow_column_letters)


def _style_from_cell(cell):
    return {
        "font": copy(cell.font) if cell.font else Font(),
        "border": copy(cell.border) if cell.border else Border(),
        "fill": copy(cell.fill) if cell.fill else PatternFill(),
        "number_format": cell.number_format,
        "alignment": copy(cell.alignment) if cell.alignment else Alignment(),
    }


def _apply_style(cell, style):
    cell.font = copy(style["font"])
    cell.border = copy(style["border"])
    cell.fill = copy(style["fill"])
    cell.number_format = style["number_format"]
    cell.alignment = copy(style["alignment"])


def _is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def _select_columns(df, columns):
    if columns is None or columns == ["*"] or columns == "*" or columns == []:
        return df

    available = [c for c in columns if c in df.columns]
    if not available:
        raise ValueError(f"指定列在 DataFrame 中不存在: {columns}")
    return df[available]


def _is_missing_scalar(value):
    if value is None:
        return True
    try:
        missing = pd.isna(value)
        if isinstance(missing, (bool, int)):
            return bool(missing)
    except Exception:
        pass
    return False


def _normalize_field_name(value):
    if _is_missing_scalar(value):
        return ""
    return str(value).strip()


def _normalize_match_value(value):
    """Normalize Excel cell values and DataFrame values for key matching."""
    if _is_missing_scalar(value):
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Real) and not isinstance(value, bool):
        number = float(value)
        if number.is_integer():
            return str(int(number))
    return str(value).strip()


def _safe_cell_value(value):
    if _is_missing_scalar(value):
        return ""
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    try:
        if hasattr(value, "item"):
            return value.item()
    except Exception:
        pass
    return value


def _series_values_for_match(df, defs):
    return [df[item["df_col"]].to_numpy(copy=False) for item in defs]


def _match_key_from_arrays(arrays, row_pos):
    return tuple(_normalize_match_value(values[row_pos]) for values in arrays)


def _build_match_row_index(df, match_defs):
    arrays = _series_values_for_match(df, match_defs)
    first_index = {}
    duplicate_keys = set()
    for row_pos in range(len(df)):
        key = _match_key_from_arrays(arrays, row_pos)
        if not any(key):
            continue
        if key in first_index:
            duplicate_keys.add(key)
            continue
        first_index[key] = row_pos
    return first_index, duplicate_keys


def _write_value_arrays(df, write_defs):
    return {
        mapping["df_col"]: df[mapping["df_col"]].to_numpy(copy=False)
        for mapping in write_defs
    }


def _build_header_map(ws, header_row, start_col, end_col):
    header_map = {}
    for col_idx in range(start_col, end_col + 1):
        field_name = _normalize_field_name(ws.cell(row=header_row, column=col_idx).value)
        if field_name and field_name not in header_map:
            header_map[field_name] = col_idx
    return header_map


def _mapping_template_field(mapping):
    return _normalize_field_name(
        mapping.get("template_field")
        or mapping.get("template_header")
        or mapping.get("template_ref")
        or mapping.get("field")
    )


def _mapping_template_col(mapping):
    return mapping.get("template_col") or mapping.get("template_column") or mapping.get("col")


def _clean_mapping_rows(rows):
    clean_rows = []
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        template_field = _mapping_template_field(row)
        template_col = _mapping_template_col(row)
        df_col = _normalize_field_name(row.get("df_col") or row.get("data_field"))
        if template_field or template_col or df_col:
            clean_rows.append(row)
    return clean_rows


def _collect_required_header_fields(payload):
    required = []
    for row in (payload.get("match_keys") or []) + (payload.get("write_mappings") or []):
        if not isinstance(row, dict):
            continue
        field = _mapping_template_field(row)
        if field and field not in required:
            required.append(field)
    return required


def _find_header_row(ws, required_fields, start_row, end_row, start_col, end_col):
    required = [_normalize_field_name(field) for field in required_fields if _normalize_field_name(field)]
    if not required:
        raise ValueError("自动查找表头需要至少一个模板字段")

    best_row = None
    best_found = []
    for row_idx in range(start_row, end_row + 1):
        header_map = _build_header_map(ws, row_idx, start_col, end_col)
        found = [field for field in required if field in header_map]
        if len(found) == len(required):
            return row_idx, header_map
        if len(found) > len(best_found):
            best_row = row_idx
            best_found = found

    missing = [field for field in required if field not in best_found]
    detail = f"，最接近的是第 {best_row} 行，找到: {', '.join(best_found)}" if best_row else ""
    raise ValueError(
        f"自动查找表头失败：在第 {start_row}-{end_row} 行未找到字段: "
        f"{', '.join(missing or required)}{detail}"
    )


def _resolve_df_col(df, value, label):
    target = _normalize_field_name(value)
    if not target:
        raise ValueError(f"{label}的数据字段不能为空")
    for column in df.columns:
        if column == value or _normalize_field_name(column) == target:
            return column
    raise ValueError(f"{label}的数据字段 [{target}] 在输入表中不存在")


def _resolve_template_col(mapping, header_map, max_row, max_col, label):
    col_ref = _mapping_template_col(mapping)
    if not _is_blank(col_ref):
        col_idx = _eval_position(col_ref, max_row, max_col, f"{label}模板列", allow_column_letters=True)
        if col_idx < 1:
            raise ValueError(f"{label}模板列必须大于 0")
        return col_idx

    field = _mapping_template_field(mapping)
    if not field:
        raise ValueError(f"{label}的模板字段/列不能为空")
    if not header_map:
        raise ValueError(f"{label}使用模板字段 [{field}]，但当前模式没有可用表头")
    if field not in header_map:
        raise ValueError(f"{label}模板字段 [{field}] 在表头中不存在")
    return header_map[field]


def _resolve_match_fill_header(ws, payload):
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1
    start_col = _eval_optional_position(
        payload.get("start_col"), max_row, max_col, "起始列", allow_column_letters=True
    ) or 1
    end_col = _eval_optional_position(
        payload.get("end_col"), max_row, max_col, "结束列", allow_column_letters=True
    ) or max_col
    start_col = max(1, start_col)
    end_col = max(1, end_col)
    if end_col < start_col:
        raise ValueError(f"结束列不能小于起始列: {end_col} < {start_col}")

    header_mode = str(payload.get("header_mode") or "specified").strip() or "specified"
    if header_mode == "manual_columns":
        return None, {}, start_col, end_col

    if header_mode == "auto_find":
        search_start = _eval_optional_position(
            payload.get("header_search_start_row"), max_row, max_col, "表头查找起始行"
        ) or 1
        search_end = _eval_optional_position(
            payload.get("header_search_end_row"), max_row, max_col, "表头查找结束行"
        ) or min(max_row, 50)
        search_start = max(1, search_start)
        search_end = min(max_row, max(1, search_end))
        if search_end < search_start:
            raise ValueError(f"表头查找结束行不能小于起始行: {search_end} < {search_start}")
        header_row, header_map = _find_header_row(
            ws,
            _collect_required_header_fields(payload),
            search_start,
            search_end,
            start_col,
            end_col,
        )
        return header_row, header_map, start_col, end_col

    header_row = _eval_position(payload.get("header_row"), max_row, max_col, "表头行")
    if header_row < 1:
        raise ValueError("表头行必须大于 0")
    return header_row, _build_header_map(ws, header_row, start_col, end_col), start_col, end_col


def _resolve_match_fill_mappings(df, header_map, payload, max_row, max_col):
    match_rows = _clean_mapping_rows(payload.get("match_keys") or [])
    if not match_rows:
        raise ValueError("匹配填充需要至少一个匹配键")

    match_defs = []
    match_template_fields = set()
    match_df_fields = set()
    for index, row in enumerate(match_rows, start=1):
        df_col = _resolve_df_col(df, row.get("df_col") or row.get("data_field"), f"匹配键 {index}")
        template_col = _resolve_template_col(row, header_map, max_row, max_col, f"匹配键 {index}")
        template_field = _mapping_template_field(row)
        if template_field:
            match_template_fields.add(template_field)
        match_df_fields.add(_normalize_field_name(df_col))
        match_defs.append({"template_col": template_col, "df_col": df_col})

    write_rows = _clean_mapping_rows(payload.get("write_mappings") or [])
    write_defs = []
    write_mode = str(payload.get("write_mode") or "auto_same_name")
    if write_mode == "manual":
        if not write_rows:
            raise ValueError("手动写入映射模式下需要至少一条写入映射")
        for index, row in enumerate(write_rows, start=1):
            df_col = _resolve_df_col(df, row.get("df_col") or row.get("data_field"), f"写入映射 {index}")
            template_col = _resolve_template_col(row, header_map, max_row, max_col, f"写入映射 {index}")
            write_defs.append({"template_col": template_col, "df_col": df_col})
    else:
        if str(payload.get("header_mode") or "specified") == "manual_columns":
            raise ValueError("手动指定列模式下必须填写写入映射")
        df_by_name = {_normalize_field_name(column): column for column in df.columns}
        for field, template_col in header_map.items():
            if field in match_template_fields or field in match_df_fields:
                continue
            if field in df_by_name:
                write_defs.append({"template_col": template_col, "df_col": df_by_name[field]})

    if not write_defs:
        raise ValueError("没有可写入映射；请填写写入映射，或确认模板表头与数据列同名")
    return match_defs, write_defs


def match_fill_template(wb, df, payload):
    """按模板已有行的匹配键查找 DataFrame 行，并把字段写入模板指定列。"""
    if openpyxl is None:
        return False, "缺少 openpyxl 依赖"
    if not isinstance(df, pd.DataFrame):
        return False, "匹配填充需要 DataFrame 输入"
    if df.empty:
        return False, "匹配填充的数据表为空"

    sheet_name = str(payload.get("sheet_name") or "").strip()
    if not sheet_name:
        sheet_name = wb.sheetnames[0] if wb.sheetnames else ""
    if sheet_name not in wb.sheetnames:
        return False, f"工作表 [{sheet_name}] 不存在"

    ws = wb[sheet_name]
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1

    try:
        header_row, header_map, _, _ = _resolve_match_fill_header(ws, payload)
        default_start = (header_row + 1) if header_row else 1
        data_start = _eval_optional_position(
            payload.get("data_start_row"), max_row, max_col, "数据起始行"
        ) or default_start
        data_end = _eval_optional_position(
            payload.get("data_end_row"), max_row, max_col, "数据结束行"
        ) or max_row
        data_start = max(1, data_start)
        data_end = max(1, data_end)
        if data_end < data_start:
            return False, f"数据结束行不能小于起始行: {data_end} < {data_start}"
        match_defs, write_defs = _resolve_match_fill_mappings(
            df, header_map, payload, max_row, max_col
        )
    except ValueError as exc:
        return False, str(exc)

    df_index, duplicate_keys = _build_match_row_index(df, match_defs)
    write_arrays = _write_value_arrays(df, write_defs)

    on_missing = str(payload.get("on_missing") or "skip")
    on_duplicate = str(payload.get("on_duplicate") or "first")
    overwrite_formulas = bool(payload.get("overwrite_formulas", False))

    matched_rows = 0
    missing_rows = 0
    duplicate_hits = 0
    blank_key_rows = 0
    protected_cells = 0
    written_cells = 0

    for row_idx in range(data_start, data_end + 1):
        template_key = tuple(
            _normalize_match_value(ws.cell(row=row_idx, column=item["template_col"]).value)
            for item in match_defs
        )
        if not any(template_key):
            blank_key_rows += 1
            continue
        row_pos = df_index.get(template_key)
        if row_pos is None:
            missing_rows += 1
            if on_missing == "error":
                return False, f"第 {row_idx} 行匹配键 {template_key} 在数据表中找不到"
            continue
        if template_key in duplicate_keys:
            duplicate_hits += 1
            if on_duplicate == "error":
                return False, f"第 {row_idx} 行匹配键 {template_key} 在数据表中匹配到多行"

        matched_rows += 1
        for mapping in write_defs:
            cell = ws.cell(row=row_idx, column=mapping["template_col"])
            if not overwrite_formulas and _is_formula(cell.value):
                protected_cells += 1
                continue
            cell.value = _safe_cell_value(write_arrays[mapping["df_col"]][row_pos])
            written_cells += 1

    msg = f"匹配填充完成：匹配 {matched_rows} 行，写入 {written_cells} 个单元格"
    notes = []
    if missing_rows:
        notes.append(f"未匹配 {missing_rows} 行")
    if duplicate_hits:
        notes.append(f"重复键取第一条 {duplicate_hits} 次")
    if blank_key_rows:
        notes.append(f"空匹配键跳过 {blank_key_rows} 行")
    if protected_cells:
        notes.append(f"公式单元格已保护 {protected_cells} 个")
    if notes:
        msg += "；" + "；".join(notes)
    return True, msg


def apply_template_insert_payload(wb, payload):
    """Apply one insert payload to a workbook, routing by insert mode."""
    if not isinstance(payload, dict) or not payload.get("_template_insert"):
        return False, "插入区域输入无效"
    if payload.get("mode") == "match_fill":
        return match_fill_template(wb, payload.get("df"), payload)

    sheet_name = payload.get("sheet_name") or (wb.sheetnames[0] if wb.sheetnames else "")
    return insert_into_template(
        wb,
        payload["df"],
        sheet_name,
        payload.get("start_row") or "max_row + 1",
        payload.get("start_col") or "1",
        columns=["*"],
        write_header=payload.get("write_header", True),
        inherit_style=False,
        end_row_expr=payload.get("end_row") or None,
        end_col_expr=payload.get("end_col") or None,
    )


def apply_template_cell_edits(wb, edits):
    """Apply explicit cell value edits to an openpyxl workbook."""
    if openpyxl is None:
        raise ImportError("缺少 openpyxl 依赖，请执行: pip install openpyxl")
    applied = []
    for index, edit in enumerate(edits or [], start=1):
        if not isinstance(edit, dict):
            continue
        sheet_name = str(edit.get("sheet_name") or "").strip()
        if not sheet_name:
            sheet_name = wb.sheetnames[0] if wb.sheetnames else ""
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"单元格编辑第 {index} 条失败：工作表 [{sheet_name}] 不存在")
        ws = wb[sheet_name]
        row_expr = edit.get("row", edit.get("row_index", ""))
        col_expr = edit.get("col", edit.get("col_index", ""))
        if _is_blank(row_expr) or _is_blank(col_expr):
            continue
        max_row = ws.max_row or 1
        max_col = ws.max_column or 1
        row = _eval_position(row_expr, max_row, max_col, f"单元格编辑第 {index} 条行")
        col = _eval_position(
            col_expr,
            max_row,
            max_col,
            f"单元格编辑第 {index} 条列",
            allow_column_letters=True,
        )
        if row < 1 or col < 1:
            raise ValueError(f"单元格编辑第 {index} 条失败：行列必须大于 0")
        value = edit.get("value", "")
        if value is None:
            value = ""
        ws.cell(row=row, column=col).value = value
        applied.append({"sheet_name": sheet_name, "row": row, "col": col, "value": value})
    return applied


def insert_into_template(
    wb,
    df,
    sheet_name,
    start_row_expr,
    start_col_expr,
    columns=None,
    write_header=True,
    inherit_style=False,
    end_row_expr=None,
    end_col_expr=None,
):
    """
    将 DataFrame 写入到模板工作表的指定区域。

    参数:
        wb:              openpyxl Workbook 对象
        df:              要写入的 DataFrame
        sheet_name:      目标工作表名称
        start_row_expr:  起始行表达式，如 "max_row+1"、"5"
        start_col_expr:  起始列表达式，如 "1"、"A"、"max_col+1"
        columns:         要写入的列列表，None、[] 或 ["*"] 表示全部列
        write_header:    是否写入表头
        inherit_style:   是否继承上一行同列的样式，默认关闭以保留模板目标区域原样式
        end_row_expr:    可选，结束行表达式；填写后会限制写入高度
        end_col_expr:    可选，结束列表达式；填写后会限制写入宽度

    返回:
        (success, message)
    """
    if openpyxl is None:
        return False, "缺少 openpyxl 依赖"

    if sheet_name not in wb.sheetnames:
        return False, f"工作表 [{sheet_name}] 不存在"

    ws = wb[sheet_name]
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1

    try:
        start_row = _eval_position(start_row_expr, max_row, max_col, "起始行")
        start_col = _eval_position(
            start_col_expr, max_row, max_col, "起始列", allow_column_letters=True
        )
        end_row = _eval_optional_position(end_row_expr, max_row, max_col, "结束行")
        end_col = _eval_optional_position(
            end_col_expr, max_row, max_col, "结束列", allow_column_letters=True
        )
        fill_df = _select_columns(df, columns)
    except ValueError as exc:
        return False, str(exc)

    start_row = max(1, start_row)
    start_col = max(1, start_col)

    if end_row is not None:
        end_row = max(1, end_row)
        if end_row < start_row:
            return False, f"结束行不能小于起始行: {end_row} < {start_row}"
    if end_col is not None:
        end_col = max(1, end_col)
        if end_col < start_col:
            return False, f"结束列不能小于起始列: {end_col} < {start_col}"

    notes = []

    if end_col is not None:
        max_write_cols = end_col - start_col + 1
        if max_write_cols <= 0:
            return False, "可写入列数为 0，请检查起止列"
        if fill_df.shape[1] > max_write_cols:
            fill_df = fill_df.iloc[:, :max_write_cols]
            notes.append(f"数据列已按结束列截断为 {max_write_cols} 列")

    if fill_df.shape[1] == 0:
        return False, "没有可写入的列"

    if end_row is not None:
        max_area_rows = end_row - start_row + 1
        header_rows = 1 if write_header else 0
        max_data_rows = max(0, max_area_rows - header_rows)
        if len(fill_df) > max_data_rows:
            fill_df = fill_df.iloc[:max_data_rows, :]
            notes.append(f"数据行已按结束行截断为 {max_data_rows} 行")

    num_rows = len(fill_df)
    num_cols = len(fill_df.columns)
    row_offset = 1 if write_header else 0
    total_rows = row_offset + num_rows

    styles = []
    if inherit_style:
        ref_row = max(1, start_row - 1)
        for j in range(num_cols):
            styles.append(_style_from_cell(ws.cell(row=ref_row, column=start_col + j)))

    protected = set()
    for r in range(start_row, start_row + total_rows):
        for c in range(start_col, start_col + num_cols):
            if _is_formula(ws.cell(row=r, column=c).value):
                protected.add((r, c))

    if write_header:
        for j, col_name in enumerate(fill_df.columns):
            row_idx = start_row
            col_idx = start_col + j
            if (row_idx, col_idx) in protected:
                continue
            cell = ws.cell(row=row_idx, column=col_idx, value=str(col_name))
            if inherit_style and j < len(styles):
                _apply_style(cell, styles[j])

    for i in range(num_rows):
        row_idx = start_row + row_offset + i
        for j in range(num_cols):
            col_idx = start_col + j
            if (row_idx, col_idx) in protected:
                continue
            val = fill_df.iat[i, j]
            if pd.isna(val):
                val = ""
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if inherit_style and j < len(styles):
                _apply_style(cell, styles[j])

    if end_row is not None or end_col is not None:
        area = (
            f"范围: 行{start_row}-{end_row if end_row is not None else start_row + total_rows - 1}, "
            f"列{start_col}-{end_col if end_col is not None else start_col + num_cols - 1}"
        )
    else:
        area = f"起始: 行{start_row}, 列{start_col}"

    msg = f"已写入 {num_rows} 行 x {num_cols} 列 到 {sheet_name} ({area})"
    if notes:
        msg += "；" + "；".join(notes)
    return True, msg


def save_template(wb, output_path):
    """保存模板到文件。"""
    out_dir = os.path.dirname(output_path)
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)
    wb.save(output_path)
    return True, output_path
